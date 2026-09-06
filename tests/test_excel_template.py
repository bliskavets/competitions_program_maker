import io

from openpyxl import load_workbook

from app.models import AgeCategory, Competition, User, WeightCategory
from app.services.excel_template import render_import_template


def test_template_uses_weight_dropdowns_and_numeric_children_weight():
    owner = User(login="owner", email="owner@example.pl", password_hash="x")
    competition = Competition(name="Puchar", owner=owner)
    cadets = AgeCategory(name="Kadeci", competition=competition)
    children = AgeCategory(name="Dzieci", competition=competition)
    WeightCategory(name="45 kg", weight=45, age_category=cadets)
    WeightCategory(name="50 kg", weight=50, age_category=cadets)

    wb = load_workbook(io.BytesIO(render_import_template(competition)))
    assert {"Instrukcja", "Kadeci", "Dzieci", "_Słowniki"}.issubset(wb.sheetnames)
    assert wb["_Słowniki"].sheet_state == "hidden"
    assert wb["Kadeci"]["D1"].value == "Actual weight"

    cadet_validations = list(wb["Kadeci"].data_validations.dataValidation)
    assert any(item.type == "list" and "C2:C500" in str(item.sqref) for item in cadet_validations)
    assert any(item.type == "decimal" and "D2:D500" in str(item.sqref) for item in cadet_validations)
    child_validations = list(wb["Dzieci"].data_validations.dataValidation)
    assert any(item.type == "decimal" and "C2:C500" in str(item.sqref) for item in child_validations)
