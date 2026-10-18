import io

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import (
    AgeCategory,
    Bout,
    Competition,
    Participant,
    Round,
    User,
    WeightCategory,
)
from tests.conftest import register


def _online_round(client, db, count=3):
    register(client, "owner", "owner@example.pl")
    owner = db.scalar(select(User).where(User.login == "owner"))
    competition = Competition(name="Puchar", owner_id=owner.id)
    age = AgeCategory(name="Dzieci", competition=competition)
    weight = WeightCategory(name="Grupa 1", age_category=age)
    for index in range(1, count + 1):
        Participant(name=f"Zawodnik {index}", order_index=index, weight_category=weight)
    db.add(competition)
    db.commit()
    response = client.post(f"/weight-categories/{weight.id}/rounds", follow_redirects=False)
    assert response.status_code == 303
    db.expire_all()
    rnd = db.scalar(select(Round).where(Round.weight_category_id == weight.id))
    return rnd


def test_online_winners_complete_standings_and_result_exports(client, db_session):
    rnd = _online_round(client, db_session)
    bouts = list(
        db_session.scalars(
            select(Bout).where(Bout.round_id == rnd.id).order_by(Bout.sequence)
        )
    )
    assert len(bouts) == 3
    for bout in bouts:
        response = client.post(
            f"/bouts/{bout.id}/winner",
            data={"winner_id": min(bout.player_a_id, bout.player_b_id)},
        )
        assert response.status_code == 200

    page = client.get(f"/rounds/{rnd.id}")
    assert "Klasyfikacja końcowa" in page.text
    assert "Wyniki Excel" in page.text

    excel = client.get(f"/rounds/{rnd.id}/results/excel")
    assert excel.status_code == 200
    workbook = load_workbook(io.BytesIO(excel.content))
    assert workbook.active.title == "Wyniki"
    assert workbook.active["A3"].value == "M-ce"

    pdf = client.get(f"/rounds/{rnd.id}/results/pdf")
    assert pdf.status_code == 200
    assert pdf.content[:4] == b"%PDF"

    combined_excel = client.get(
        f"/competitions/{rnd.weight_category.age_category.competition_id}/results/excel"
    )
    assert combined_excel.status_code == 200
    assert load_workbook(io.BytesIO(combined_excel.content)).active.title == "Wyniki zawodów"
    combined_pdf = client.get(
        f"/competitions/{rnd.weight_category.age_category.competition_id}/results/pdf"
    )
    assert combined_pdf.status_code == 200
    assert combined_pdf.content[:4] == b"%PDF"


def test_online_result_rejects_outsider(client, db_session):
    rnd = _online_round(client, db_session)
    bout = db_session.scalar(select(Bout).where(Bout.round_id == rnd.id))
    response = client.post(f"/bouts/{bout.id}/winner", data={"winner_id": 999999})
    assert response.status_code == 400


def test_online_cyclic_tie_can_start_replay_without_manual_names(client, db_session):
    rnd = _online_round(client, db_session, count=4)
    participants = list(
        db_session.scalars(
            select(Participant)
            .where(Participant.weight_category_id == rnd.weight_category_id)
            .order_by(Participant.order_index)
        )
    )
    bouts = list(
        db_session.scalars(
            select(Bout).where(Bout.round_id == rnd.id, Bout.stage == "group")
        )
    )
    leader, second, third, fourth = participants
    outcomes = {
        frozenset((leader.id, second.id)): leader.id,
        frozenset((leader.id, third.id)): leader.id,
        frozenset((leader.id, fourth.id)): leader.id,
        frozenset((second.id, third.id)): second.id,
        frozenset((third.id, fourth.id)): third.id,
        frozenset((fourth.id, second.id)): fourth.id,
    }
    for bout in bouts:
        response = client.post(
            f"/bouts/{bout.id}/winner",
            data={"winner_id": outcomes[frozenset((bout.player_a_id, bout.player_b_id))]},
        )
        assert response.status_code == 200

    page = client.get(f"/rounds/{rnd.id}")
    assert "Powtórz walki" in page.text
    assert "Klasyfikacja końcowa" not in page.text

    response = client.post(
        f"/rounds/{rnd.id}/tiebreak",
        data={
            "group_name": "S",
            "participant_id": [str(second.id), str(third.id), str(fourth.id)],
        },
    )
    assert response.status_code == 200
    assert "Dogrywka S" in response.text
    replays = list(
        db_session.scalars(
            select(Bout).where(Bout.round_id == rnd.id, Bout.stage == "tiebreak")
        )
    )
    assert len(replays) == 3
    assert {replay.player_a_id for replay in replays} | {
        replay.player_b_id for replay in replays
    } == {second.id, third.id, fourth.id}
