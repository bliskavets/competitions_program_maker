"""Generate a competition-aware workbook for participant imports."""
from __future__ import annotations

import io
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, quote_sheetname
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from app.models import Competition
from app.services.child_grouping import is_children_category

HEADERS = [
    "Name and Surname",
    "Year of Birth",
    "Weight category",
    "Actual weight",
    "Team",
    "Other Information",
]
DEFAULT_AGES = ["Dzieci", "Młodzicy", "Kadeci", "Juniorzy"]


def render_import_template(competition: Competition) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    dictionary = wb.create_sheet("_Słowniki")
    dictionary.sheet_state = "hidden"

    ages = list(competition.age_categories)
    sheet_specs = [
        (age.name or f"Kategoria {index}", [wc.name for wc in age.weight_categories if wc.name])
        for index, age in enumerate(ages, start=1)
    ] or [(name, []) for name in DEFAULT_AGES]

    used_titles: set[str] = set()
    for index, (age_name, weights) in enumerate(sheet_specs, start=1):
        title = _unique_title(age_name, used_titles)
        ws = wb.create_sheet(title)
        ws.append(HEADERS)
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8E8E8")
        widths = [28, 16, 20, 16, 24, 30]
        for column, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(column)].width = width

        if is_children_category(age_name):
            validation = DataValidation(
                type="decimal",
                operator="between",
                formula1="1",
                formula2="300",
                allow_blank=True,
            )
            validation.promptTitle = "Waga rzeczywista"
            validation.prompt = "Wpisz zmierzoną wagę w kg, np. 24,50."
            validation.error = "Waga musi być liczbą od 1 do 300 kg."
            validation.errorTitle = "Nieprawidłowa waga"
            validation.showErrorMessage = True
            validation.showInputMessage = True
            ws.add_data_validation(validation)
            validation.add("C2:C500")
        elif weights:
            column = index
            for row, value in enumerate(weights, start=1):
                dictionary.cell(row=row, column=column, value=value)
            col_letter = get_column_letter(column)
            range_name = f"weight_categories_{index}"
            reference = f"{quote_sheetname(dictionary.title)}!${col_letter}$1:${col_letter}${len(weights)}"
            wb.defined_names.add(DefinedName(range_name, attr_text=reference))
            validation = DataValidation(type="list", formula1=range_name, allow_blank=False)
            validation.error = "Wybierz kategorię z listy."
            validation.errorTitle = "Nieznana kategoria"
            validation.showErrorMessage = True
            ws.add_data_validation(validation)
            validation.add("C2:C500")

        measured = DataValidation(
            type="decimal", operator="between", formula1="1", formula2="300", allow_blank=True
        )
        measured.error = "Waga musi być liczbą od 1 do 300 kg."
        measured.showErrorMessage = True
        ws.add_data_validation(measured)
        measured.add("D2:D500")

    info = wb.create_sheet("Instrukcja", 0)
    info["A1"] = "Import uczestników"
    info["A1"].font = Font(bold=True, size=14)
    info["A3"] = "Każdy arkusz odpowiada kategorii wiekowej. Nie zmieniaj nazw nagłówków."
    info["A4"] = "Dzieci: w kolumnie Weight category wpisz faktyczną wagę; grupy powstaną automatycznie."
    info["A5"] = "Pozostałe kategorie: wybierz kategorię wagową z listy, jeśli została skonfigurowana."
    info.column_dimensions["A"].width = 105

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _unique_title(value: str, used: set[str]) -> str:
    base = re.sub(r"[\\/*?:\[\]]", "-", value).strip()[:31] or "Kategoria"
    title = base
    suffix = 2
    while title.lower() in used:
        marker = f" {suffix}"
        title = f"{base[: 31 - len(marker)]}{marker}"
        suffix += 1
    used.add(title.lower())
    return title
